"""Произвольные расходы (удержания) перевозчика — CarrierAdjustment (2026-09-30).

Проверяем: CRUD и права, учёт в балансе (вычитается ДО %СК, как штраф),
шаблоны описаний, отдельные строки + формулы в недельной выгрузке.
"""
from datetime import date, datetime
from io import BytesIO

import openpyxl
from sqlmodel import Session

from app import models
from .conftest import make_user, token_headers

CARRIER = "ТестПеревозчикУдерж"
WEEK = date(2026, 9, 7)          # понедельник
BASE = "/api/carriers/balance"


def _seed(db):
    with Session(db) as s:
        s.add(models.Carrier(name=CARRIER, insurance_pct=10))
        s.add(models.Trip(request_number="ADJ-1", carrier_name=CARRIER, status="Завершено",
                          dep_at=datetime(2026, 9, 8, 10, 0), amount=10000, fines=1000,
                          report_week=WEEK))
        s.commit()


def _row(client, headers):
    rows = client.get(f"{BASE}/", headers=headers).json()
    return next(r for r in rows if r["carrier_name"] == CARRIER)


def test_adjustment_flow(client, db, admin_headers):
    _seed(db)
    before = _row(client, admin_headers)
    assert before["net"] == 8100.0                    # (10000-1000)*0.9

    # Любая дата недели снапается на понедельник; описание сохраняем в шаблоны.
    r = client.post(f"{BASE}/adjustments", headers=admin_headers, json={
        "carrier_name": CARRIER, "report_week": "2026-09-10", "amount": 2000,
        "description": "ГСМ", "save_preset": True})
    assert r.status_code == 200, r.text
    adj = r.json()
    assert adj["report_week"] == "2026-09-07"

    after = _row(client, admin_headers)
    assert after["adjustments"] == 2000.0
    assert after["net"] == 6300.0                     # (10000-1000-2000)*0.9 — до СК
    wk = next(w for w in after["weeks"] if w["week_start"] == "2026-09-07")
    assert wk["adjustments"] == 2000.0

    presets = client.get(f"{BASE}/expense-presets", headers=admin_headers).json()
    assert any(p["text"] == "ГСМ" for p in presets)

    lst = client.get(f"{BASE}/adjustments", params={"carrier": CARRIER}, headers=admin_headers).json()
    assert [a["id"] for a in lst] == [adj["id"]]

    # Выгрузка: строка «Расход: …» с суммой в N, формулы итогов учитывают расходы.
    x = client.get(f"{BASE}/weekly-export", params={"carrier": CARRIER}, headers=admin_headers)
    assert x.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(x.content))
    ws = next(wb[n] for n in wb.sheetnames if n != "Сводная")
    assert ws.cell(row=1, column=14).value == "Расход (удержание)"
    notes = [(ws.cell(row=i, column=13).value, ws.cell(row=i, column=14).value) for i in range(2, ws.max_row + 1)]
    assert ("Расход: ГСМ", 2000) in notes
    assert ws["P9"].value.startswith("=SUM($N$2:$N$")
    assert ws["P10"].value == "=P7-P8-P9"
    sv = wb["Сводная"]
    assert sv.cell(row=5, column=9).value == "Расходы, ₽"
    assert sv.cell(row=6, column=10).value == "=G6-H6-I6"

    # Правка и удаление.
    r = client.put(f"{BASE}/adjustments/{adj['id']}", headers=admin_headers, json={
        "carrier_name": CARRIER, "report_week": "2026-09-07", "amount": 500, "description": "ГСМ"})
    assert r.status_code == 200 and r.json()["amount"] == 500
    assert client.delete(f"{BASE}/adjustments/{adj['id']}", headers=admin_headers).status_code == 200
    assert _row(client, admin_headers)["net"] == 8100.0


def test_adjustment_validation_and_roles(client, db, admin_headers):
    bad = {"carrier_name": CARRIER, "report_week": "2026-09-07", "amount": 0, "description": "x"}
    assert client.post(f"{BASE}/adjustments", headers=admin_headers, json=bad).status_code == 400
    bad = {"carrier_name": CARRIER, "report_week": "2026-09-07", "amount": 10, "description": "  "}
    assert client.post(f"{BASE}/adjustments", headers=admin_headers, json=bad).status_code == 400
    bad = {"carrier_name": "НетТакого", "report_week": "2026-09-07", "amount": 10, "description": "x"}
    assert client.post(f"{BASE}/adjustments", headers=admin_headers, json=bad).status_code == 404

    foreman = make_user(db, "adj_foreman", "foreman")
    ok = {"carrier_name": CARRIER, "report_week": "2026-09-07", "amount": 10, "description": "x"}
    assert client.post(f"{BASE}/adjustments", headers=token_headers(foreman), json=ok).status_code == 403
    # чтение staff разрешено
    assert client.get(f"{BASE}/adjustments", params={"carrier": CARRIER},
                      headers=token_headers(foreman)).status_code == 200


def test_balance_uses_reporting_weeks(client, db, admin_headers):
    name = "ТестНеделиУчёта"
    with Session(db) as s:
        s.add(models.Carrier(name=name, insurance_pct=0))
        # Отгрузка в неделю 07.09, учёт рейса — 14.09, штраф поступил к учёту 21.09.
        s.add(models.Trip(request_number="RW-1", carrier_name=name, status="Завершено",
                          dep_at=datetime(2026, 9, 9, 10, 0), amount=5000, fines=300,
                          report_week=date(2026, 9, 14), fines_report_week=date(2026, 9, 21)))
        # Старый рейс без недели учёта — фолбэк на неделю отгрузки (07.09).
        s.add(models.Trip(request_number="RW-2", carrier_name=name, status="Завершено",
                          dep_at=datetime(2026, 9, 8, 10, 0), amount=1000, fines=0))
        s.commit()
    rows = client.get(f"{BASE}/", headers=admin_headers).json()
    weeks = {w["week_start"]: w for w in next(r for r in rows if r["carrier_name"] == name)["weeks"]}
    assert weeks["2026-09-07"]["gross"] == 1000 and weeks["2026-09-07"]["trips"] == 1
    assert weeks["2026-09-14"]["gross"] == 5000 and weeks["2026-09-14"]["fines"] == 0
    assert weeks["2026-09-21"]["fines"] == 300 and weeks["2026-09-21"]["trips"] == 0
